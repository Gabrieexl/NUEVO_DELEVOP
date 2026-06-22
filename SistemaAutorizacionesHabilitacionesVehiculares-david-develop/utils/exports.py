"""
Utilidades para exportación a Excel usando openpyxl.
"""

from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.utils import timezone


class ExcelExporter:
    """
    Clase para exportar datos a Excel de manera sencilla.
    """
    
    def __init__(self, titulo, headers):
        """
        Inicializa el exportador.
        
        Args:
            titulo: Título del reporte
            headers: Lista de encabezados de columna
        """
        self.titulo = titulo
        self.headers = headers
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = titulo[:31]
    
    def _aplicar_estilos_cabecera(self, celda):
        """Aplica estilos a celdas de cabecera."""
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        celda.alignment = Alignment(horizontal='center', vertical='center')
        celda.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    def _aplicar_estilos_celda(self, celda):
        """Aplica estilos a celdas de datos."""
        celda.alignment = Alignment(vertical='center')
        celda.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
    
    def export_to_response(self, data, filename):
        """
        Exporta los datos a una respuesta HTTP con archivo Excel.
        
        Args:
            data: Lista de listas con los datos
            filename: Nombre del archivo (sin extensión)
        
        Returns:
            HttpResponse con el archivo Excel
        """
        # Título del reporte
        self.ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(self.headers))
        celda_titulo = self.ws.cell(row=1, column=1, value=self.titulo)
        celda_titulo.font = Font(bold=True, size=14)
        celda_titulo.alignment = Alignment(horizontal='center')
        
        # Fecha de generación
        fecha_actual = timezone.localtime(timezone.now()).strftime('%d/%m/%Y %H:%M')
        self.ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(self.headers))
        celda_fecha = self.ws.cell(row=2, column=1, value=f'Generado el: {fecha_actual}')
        celda_fecha.font = Font(italic=True, size=10)
        celda_fecha.alignment = Alignment(horizontal='center')
        
        # Cabeceras (fila 4)
        for col_num, header in enumerate(self.headers, 1):
            celda = self.ws.cell(row=4, column=col_num, value=header)
            self._aplicar_estilos_cabecera(celda)
        
        # Datos
        for row_num, row_data in enumerate(data, 5):
            for col_num, value in enumerate(row_data, 1):
                celda = self.ws.cell(row=row_num, column=col_num, value=value)
                self._aplicar_estilos_celda(celda)
        
        # Ajustar anchos de columna
        for col_num in range(1, len(self.headers) + 1):
            column_letter = get_column_letter(col_num)
            max_length = len(self.headers[col_num - 1])
            
            for row in self.ws.iter_rows(min_row=5, min_col=col_num, max_col=col_num):
                for cell in row:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
            
            self.ws.column_dimensions[column_letter].width = min(max_length + 2, 50)
        
        # Generar respuesta
        output = BytesIO()
        self.wb.save(output)
        output.seek(0)
        
        fecha_archivo = timezone.localtime(timezone.now()).strftime('%Y%m%d_%H%M')
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}_{fecha_archivo}.xlsx"'
        
        return response


def exportar_queryset_a_excel(queryset, columnas, titulo='Exportación', nombre_archivo='exportacion'):
    """
    Función auxiliar para exportar un queryset a Excel.
    
    Args:
        queryset: QuerySet de Django con los datos a exportar.
        columnas: Lista de tuplas (campo, titulo_columna) o 
                  (campo, titulo_columna, funcion_formato).
        titulo: Título del reporte.
        nombre_archivo: Nombre del archivo sin extensión.
    
    Returns:
        HttpResponse con el archivo Excel.
    """
    headers = [col[1] if len(col) > 1 else col[0] for col in columnas]
    
    data = []
    for obj in queryset:
        row = []
        for columna in columnas:
            campo = columna[0]
            funcion_formato = columna[2] if len(columna) > 2 else None
            
            # Obtener valor (soporta campos relacionados con __)
            valor = obj
            for attr in campo.split('__'):
                if valor is not None:
                    valor = getattr(valor, attr, None)
                    if callable(valor):
                        try:
                            valor = valor()
                        except:
                            pass
            
            # Aplicar función de formato si existe
            if funcion_formato and valor is not None:
                try:
                    valor = funcion_formato(valor)
                except:
                    pass
            
            row.append(valor)
        data.append(row)
    
    exporter = ExcelExporter(titulo, headers)
    return exporter.export_to_response(data, nombre_archivo)
